import socket
import os, subprocess
from path_finder import *
import json
from datetime import datetime, date, timedelta
import openpyxl
from xlsx import xlsx_value
def is_today_holiday():
    file_path = f"{current_dir}/files/nse_holidays.json"
    today_str = datetime.now().strftime("%Y-%m-%d")

    try:
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        # JSON से सीधे 'closed_days' वाली लिस्ट निकालें
        if isinstance(data, dict):
            closed_list = data.get("closed_days", [])
        else:
            closed_list = data

        # छुट्टी है तो True, नहीं तो False
        if today_str in closed_list:
            return True
            
        return False

    except Exception:
        return False
def check_internet(host="8.8.8.8", port=53, timeout=3):
    try:
        socket.setdefaulttimeout(timeout)
        socket.socket(socket.AF_INET, socket.SOCK_STREAM).connect((host, port))
        return True
    except socket.error as ex:
        return False









def update_expiry_days():
    # 1. Closed days
    with open(f"{path}nse_holidays.json", "r", encoding="utf-8") as f:
        closed_days = set(json.load(f).get("closed_days", []))

    # 2. Expiry date from main.xlsx
    raw_exp = xlsx_value("c6", "main.xlsx")
    if isinstance(raw_exp, (datetime, date)):
        exp_date = raw_exp if isinstance(raw_exp, date) else raw_exp.date()
    else:
        exp_date = datetime.strptime(str(raw_exp).strip()[:10], "%Y-%m-%d").date()

    # 3. T1 to T7 dates
    t_days = {}
    curr = exp_date
    t_num = 1
    while t_num <= 7:
        d_str = curr.strftime("%Y-%m-%d")
        if curr.weekday() < 5 and d_str not in closed_days:
            t_days[f"T{t_num}"] = curr
            t_num += 1
        curr -= timedelta(days=1)

    # 4. Target trading date (Today or Next Open Day)
    target_dt = date.today()
    while True:
        d_str = target_dt.strftime("%Y-%m-%d")
        if target_dt.weekday() < 5 and d_str not in closed_days:
            break
        target_dt += timedelta(days=1)

    # 5. ex.xlsx ko SINGLE LOAD & SINGLE SAVE me niptao (Fast I/O)
    file_path = f"{current_dir}/files/ex.xlsx"
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    status_val = "N/A"
    cols = ["A", "B", "C", "D", "E", "F", "G"]
    
    for i, col in enumerate(cols, start=1):
        t_key = f"T{i}"
        d_val = t_days.get(t_key)
        cell = f"{col}2"
        
        if d_val:
            if d_val == target_dt:
                status_val = t_key
                ws[cell] = d_val.strftime("%Y-%m-%d")
            elif d_val < target_dt:
                ws[cell] = None  # Jo din nikal gaya wo blank
            else:
                ws[cell] = d_val.strftime("%Y-%m-%d")
        else:
            ws[cell] = None

    ws["H2"] = status_val
    wb.save(file_path)